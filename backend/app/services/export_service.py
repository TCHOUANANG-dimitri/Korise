"""Full-history export of a business (Excel + PDF), offered to the owner before deleting their
account. The Excel file holds every line; the PDF is the readable summary (months, closings,
team, stock) — built by pdf_service.history_pdf from the same data."""

import io
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.models.business import Business
from app.models.catalog import Product
from app.models.customer import Customer
from app.models.events import DailyClosing, MoneyMovement, MoneyMovementType, Sale, StockMovement
from app.models.shift import Shift
from app.models.user import User, UserRole
from app.services.pdf_service import MOVEMENT_LABELS, PAYMENT_LABELS

STOCK_LABELS = {"sale": "Vente", "restock": "Entrée de stock", "adjustment": "Ajustement"}


@dataclass
class History:
    business: Business
    users: dict[uuid.UUID, User]
    products: list[Product]
    customers: list[Customer]
    sales: list[Sale]
    money: list[MoneyMovement]
    stock: list[StockMovement]
    closings: list[DailyClosing]
    shifts: list[Shift]
    balances: dict[uuid.UUID, int] = field(default_factory=dict)

    def user_name(self, user_id: uuid.UUID | None) -> str:
        user = self.users.get(user_id) if user_id else None
        return user.full_name if user else "—"


def collect(session: Session, business_id: uuid.UUID) -> History:
    def rows(model, order):
        return session.exec(select(model).where(model.business_id == business_id).order_by(order)).all()

    history = History(
        business=session.get(Business, business_id),
        users={u.id: u for u in rows(User, User.created_at)},
        products=rows(Product, Product.name),
        customers=rows(Customer, Customer.full_name),
        sales=rows(Sale, Sale.created_at),
        money=rows(MoneyMovement, MoneyMovement.created_at),
        stock=rows(StockMovement, StockMovement.created_at),
        closings=rows(DailyClosing, DailyClosing.closing_date),
        shifts=rows(Shift, Shift.opened_at),
    )
    balances: dict[uuid.UUID, int] = defaultdict(int)
    for s in history.sales:
        if s.payment_method == "credit" and s.customer_id:
            balances[s.customer_id] += s.total_amount
    for m in history.money:
        if m.type == MoneyMovementType.credit_repayment and m.customer_id:
            balances[m.customer_id] -= m.amount
    history.balances = dict(balances)
    return history


def _user_status(user: User) -> str:
    if user.deleted_at is not None:
        return "Supprimé"
    return "Actif" if user.is_active else "Désactivé"


def _dt(value: datetime | None):
    return value.replace(tzinfo=None) if isinstance(value, datetime) else value


def history_xlsx(history: History) -> bytes:
    products = {p.id: p.name for p in history.products}
    customers = {c.id: c.full_name for c in history.customers}
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, header: list[str], lines: list[list]) -> None:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="000000")
        for line in lines:
            ws.append(line)
        ws.freeze_panes = "A2"
        for i, name in enumerate(header, start=1):
            width = max([len(str(name))] + [len(str(line[i - 1])) for line in lines[:200] if line[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 45)

    sheet(
        "Ventes",
        ["Date", "Produit", "Quantité", "Prix unitaire (FCFA)", "Total (FCFA)", "Paiement", "Client", "Vendeur"],
        [
            [_dt(s.created_at), products.get(s.product_id, "—"), s.quantity, s.unit_price, s.total_amount,
             PAYMENT_LABELS.get(s.payment_method, s.payment_method), customers.get(s.customer_id, ""), history.user_name(s.user_id)]
            for s in history.sales
        ],
    )
    sheet(
        "Argent",
        ["Date", "Type", "Canal", "Montant (FCFA)", "Catégorie", "Motif", "Client", "Utilisateur"],
        [
            [_dt(m.created_at), MOVEMENT_LABELS.get(m.type.value, m.type.value),
             PAYMENT_LABELS.get(m.channel.value if m.channel else "cash"), m.amount, m.category or "", m.reason or "",
             customers.get(m.customer_id, ""), history.user_name(m.user_id)]
            for m in history.money
        ],
    )
    sheet(
        "Stock",
        ["Date", "Produit", "Type", "Quantité", "Motif", "Utilisateur"],
        [
            [_dt(m.created_at), products.get(m.product_id, "—"), STOCK_LABELS.get(m.type.value, m.type.value),
             m.quantity_delta, m.reason or "", history.user_name(m.user_id)]
            for m in history.stock
        ],
    )
    sheet(
        "Clôtures",
        ["Jour", "Espèces attendues", "Espèces comptées", "Écart espèces", "MoMo attendu", "MoMo compté", "Écart MoMo",
         "Orange attendu", "Orange compté", "Écart Orange", "Note", "Par"],
        [
            [c.closing_date, c.expected_cash, c.actual_cash, c.difference, c.expected_momo, c.actual_momo, c.difference_momo,
             c.expected_orange, c.actual_orange, c.difference_orange, c.note or "", history.user_name(c.user_id)]
            for c in history.closings
        ],
    )
    sheet(
        "Shifts",
        ["Employé", "Ouverture", "Fermeture", "Fond de caisse", "Attendu", "Compté", "Écart", "Note"],
        [
            [history.user_name(sh.user_id), _dt(sh.opened_at), _dt(sh.closed_at), sh.opening_cash, sh.expected_cash,
             sh.counted_cash, sh.difference, sh.note or ""]
            for sh in history.shifts
        ],
    )
    sheet(
        "Clients",
        ["Nom", "Téléphone", "Solde dû (FCFA)"],
        [[c.full_name, c.phone or "", history.balances.get(c.id, 0)] for c in history.customers],
    )
    sheet(
        "Produits",
        ["Nom", "Catégorie", "Code-barres", "Stock", "Seuil", "Prix d'achat", "Prix de vente", "Type", "Statut"],
        [
            [p.name, p.category or "", p.barcode or "", p.quantity if p.is_stockable else "", p.minimum_stock,
             p.purchase_price, p.selling_price, "Produit" if p.is_stockable else "Service", "Actif" if p.is_active else "Archivé"]
            for p in history.products
        ],
    )
    sheet(
        "Équipe",
        ["Nom", "Rôle", "Téléphone", "Statut", "Créé le"],
        [
            [u.full_name, "Propriétaire" if u.role == UserRole.owner else "Employé", u.phone or "", _user_status(u), _dt(u.created_at)]
            for u in history.users.values()
        ],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
